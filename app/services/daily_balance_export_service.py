from __future__ import annotations

import csv
import textwrap
from datetime import date, datetime
from decimal import Decimal
from io import BytesIO, StringIO
from math import ceil
from numbers import Real

from flask import Response, send_file
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill

from app.services.export_service import ExportService
from app.utils.dates import manila_date


class DailyBalanceExportService(ExportService):
    """Readable exports of the same live ledger and saved cash snapshots."""

    PDF_TEMPLATE = "admin/daily_balance_report_pdf.html"
    METHODS = ("cash", "gcash", "bdo", "bpi", "queenbank", "unclassified")
    FLOW_LABELS = {
        "checkout": "Checkout sales", "credit_applied": "Customer credit applied",
        "credit_received": "Customer credit received", "collection": "Debt collected",
        "adjustment_income": "Other income", "owner_topup": "Owner cash added",
        "cash_withdrawal": "Cash taken out", "expense": "Business expenses",
        "external_expense": "External-funded expenses", "payable": "Business debt paid",
        "external_payable": "External-funded debt paid", "adjustment_expense": "Budget spent",
        "refund": "Refunds paid", "owner_reimbursement": "Owner repaid (part of debt paid)",
        "expense_paid": "Expense payments before reversals", "expense_void": "Expense reversals",
        "payable_paid": "All debt payments",
    }
    FLOWS = tuple(FLOW_LABELS)
    REPORT_GROUPS = {
        "Sales": (("Cowork Bill", "cowork_bill"), ("Cafe Bill", "cafe_bill"),
                  ("Discounts", "discounts_total"), ("Cowork Discounts", "cowork_discounts"),
                  ("Cafe Discounts", "cafe_discounts"), ("Checkout Sales", "total_revenue"),
                  ("Refunds Paid", "total_refunds"),
                  ("Sales Reconciliation Difference", "sales_reconciliation_difference")),
        "Other money received": (("Debt Collected", "total_collections"),
                                 ("Other Income", "total_other_income"), ("Owner Cash Added", "total_owner_topups")),
        "Spending and transfers": (("Total Expenses", "total_expenses"),
                                   ("Business Expenses", "total_business_expenses"),
                                   ("External-funded Expenses", "total_external_expenses"),
                                   ("Debt Paid", "total_payables_paid"), ("Business Debt Paid", "total_business_payables_paid"),
                                   ("External-funded Debt Paid", "total_external_payables_paid"),
                                   ("Supplier Payments", "total_supplier_payments"),
                                   ("Owner Reimbursements (included in Debt Paid)", "total_owner_reimbursements"),
                                   ("Budget Spent", "total_budget_spend"), ("Cash Taken Out", "total_cash_withdrawals")),
        "Totals": (("Net Movement (all methods)", "net_balance"), ("Cash Movement", "cash_on_hand")),
        "Activity": (("Total Orders", "total_orders"), ("Timed Sessions", "total_sessions"),
                     ("Debt Payment Count", "collection_count")),
    }
    MONEY_FORMAT = '"₱"#,##0.00;[Red]("₱"#,##0.00);"₱"0.00'

    @staticmethod
    def _period_label(reports, soft_entries=()):
        days = sorted([str(r["report_date"]) for r in reports]
                      + [str(s["balance_date"]) for s in soft_entries])
        return f"{days[0]} to {days[-1]}" if days else "No recorded dates"

    @staticmethod
    def _payment_totals(reports):
        return {**{f"total_{m}": sum(r.get(f"{m}_total", 0) for r in reports)
                   for m in DailyBalanceExportService.METHODS[:-1]},
                **{f"{m}_count": sum(r.get(f"{m}_count", 0) for r in reports)
                   for m in DailyBalanceExportService.METHODS[:-1]}}

    @staticmethod
    def _method_rows(reports):
        rows = []
        for method in DailyBalanceExportService.METHODS:
            values = {flow: sum(r.get(f"{flow}_methods", {}).get(method, 0) for r in reports)
                      for flow in DailyBalanceExportService.FLOWS}
            active = any(abs(value) >= 0.005 for value in values.values())
            if method == "unclassified" and not active:
                continue
            received = (values["checkout"] - values["credit_applied"] + values["credit_received"]
                        + values["collection"] + values["adjustment_income"] + values["owner_topup"])
            paid_out = (values["cash_withdrawal"] + values["expense"] + values["payable"]
                        + values["adjustment_expense"] + values["refund"])
            rows.append({"method": method, **values, "received": received, "paid_out": paid_out,
                         "net": received - paid_out, "active": active})
        return rows

    @classmethod
    def report_sections(cls, report):
        sections = {group: [(label, report.get(key, 0)) for label, key in fields]
                    for group, fields in cls.REPORT_GROUPS.items()}
        sections["Checkout methods"] = [(method.upper(), report.get(f"{method}_total", 0)) for method in cls.METHODS
                                        if method != "unclassified" or report.get("unclassified_total", 0)]
        sections["Checkout counts"] = [(method.upper(), report.get(f"{method}_count", 0)) for method in cls.METHODS
                                       if method != "unclassified" or report.get("unclassified_count", 0)]
        boxes = report.get("cash_boxes")
        for label, box in (("Petty", "petty"), ("Customer", "customer")):
            sections[f"{label} cash"] = [
                (f"{label} {title}", boxes.get(key, {}).get(box) if boxes else None)
                for title, key in (("Opening", "opening"), ("Received", "receipts"), ("Top-ups", "topups"),
                                   ("Spending", "spending"), ("Withdrawn", "withdrawals"), ("Closing", "closing"))
            ]
        sections["Saved cash snapshot"] = [
            ("Saved Cash Movement", report.get("snapshot_cash_movement")),
            ("Saved Petty Closing", report.get("snapshot_petty_cash_closing")),
            ("Saved Customer Closing", report.get("snapshot_customer_cash_closing")),
            ("Cash Snapshot Difference", report.get("cash_snapshot_difference")),
            ("Petty Snapshot Difference", report.get("petty_cash_snapshot_difference")),
            ("Customer Snapshot Difference", report.get("customer_cash_snapshot_difference")),
            ("Unassigned Cash Movement", boxes.get("unassigned_cash_movement", 0) if boxes else None),
        ]
        sections["Audit"] = [("Report Status", "Live totals; saved cash snapshot" if report.get("id") is not None else "Live totals; no saved report"),
                             ("Generated By", report.get("generated_by", "Not generated")),
                             ("Generated At (UTC)", report.get("generated_at")), ("Notes", report.get("notes"))]
        return sections

    @staticmethod
    def soft_sections(entry):
        return {
            "Saved totals": [("Checkout Sales", entry.get("total_revenue", 0)),
                             ("Total Expenses", entry.get("total_expenses", 0)),
                             ("Debt Collected", entry.get("total_collections", 0)),
                             ("Other Income", entry.get("total_other_income", 0)),
                             ("Debt Paid", entry.get("total_payables_paid", 0)),
                             ("Budget Spent", entry.get("total_budget_spend", 0)),
                             ("Net Movement", entry.get("net_balance", 0))],
            "Cash snapshot": [("Saved Petty Closing", entry.get("petty_cash_closing")),
                              ("Saved Customer Closing", entry.get("customer_cash_closing")),
                              ("Saved Cash Movement", entry.get("cash_movement")),
                              ("Petty Difference Now", entry.get("petty_cash_live_difference")),
                              ("Customer Difference Now", entry.get("customer_cash_live_difference")),
                              ("Combined Difference Now", entry.get("live_cash_difference"))],
            "Audit": [("Generated By", entry.get("generated_by")),
                      ("Generated At (UTC)", entry.get("generated_at")), ("Notes", entry.get("notes"))],
        }

    def build_pdf_context(self, reports, soft_entries):
        totals = {key: sum(r.get(key, 0) for r in reports)
                  for fields in self.REPORT_GROUPS.values() for _, key in fields}
        return {
            **totals, **self._payment_totals(reports), "reports": reports, "soft_entries": soft_entries,
            "method_rows": self._method_rows(reports), "flow_labels": self.FLOW_LABELS,
            "report_details": [(r, self.report_sections(r)) for r in reports],
            "soft_details": [(s, self.soft_sections(s)) for s in soft_entries],
            "generated_at": datetime.utcnow().strftime("%Y-%m-%d %H:%M UTC"),
            "period_label": self._period_label(reports, soft_entries),
            "pdf_text_lines": self._pdf_text_lines,
            "text_parts": self._text_parts,
        }

    @staticmethod
    def _text_parts(value):
        return [value[i:i + 400] for i in range(0, len(value), 400)] if isinstance(value, str) and len(value) > 400 else [value]

    @staticmethod
    def _pdf_text_lines(value):
        # The fallback renderer does not wrap long words in table cells.
        return [part for line in str(value).splitlines()
                for part in textwrap.wrap(line, 28, replace_whitespace=False, drop_whitespace=False) or [""]]

    def export_pdf(self, reports, soft_entries):
        pdf_bytes = self.render_pdf_bytes(self.PDF_TEMPLATE, **self.build_pdf_context(reports, soft_entries))
        filename = f"daily_sales_balance_{manila_date(datetime.utcnow())}.pdf"
        return Response(pdf_bytes, mimetype="application/pdf",
                        headers={"Content-Disposition": f"attachment; filename={filename}"})

    @staticmethod
    def _csv_text(value):
        text = str(value or "")
        return "'" + text if text.startswith(("=", "+", "-", "@", "\t", "\r")) else text

    @classmethod
    def export_csv(cls, reports, soft_entries=None):
        output = StringIO()
        writer = csv.writer(output)
        writer.writerow(["Date", "Section", "Item", "Method", "Amount", "Notes"])
        writer.writerow(["", "Guide", "Amounts in PHP; Activity and Checkout counts are counts", "", "",
                         "Daily summaries, payment details, cash balances and saved snapshots overlap. Filter by Section; do not sum different sections together."])
        for report in reports:
            for section, items in cls.report_sections(report).items():
                for label, value in items:
                    numeric = isinstance(value, (Real, Decimal))
                    unit = "Count" if section in {"Activity", "Checkout counts"} else "PHP"
                    writer.writerow([report["report_date"], f"Daily summary / {section}", label, "",
                                     f"{value:.2f}" if numeric else "",
                                     unit if numeric else cls._csv_text(value)])
            for method in cls._method_rows([report]):
                for flow, label in cls.FLOW_LABELS.items():
                    writer.writerow([report["report_date"], "Payment detail", label, method["method"], f"{method[flow]:.2f}", "PHP; included in daily summaries"])
        for entry in soft_entries or []:
            for section, items in cls.soft_sections(entry).items():
                for label, value in items:
                    numeric = isinstance(value, (Real, Decimal))
                    writer.writerow([entry["balance_date"], f"Saved snapshot / {entry['period']} #{entry['id']} / {section}",
                                     label, "", f"{value:.2f}" if numeric else "",
                                     "PHP; snapshot, not additional activity" if numeric else cls._csv_text(value)])
        return send_file(BytesIO(output.getvalue().encode("utf-8-sig")), mimetype="text/csv",
                         as_attachment=True, download_name="daily_balance_details_v2.csv")

    @classmethod
    def _sheet(cls, wb, name, headers, rows, widths, money_columns=()):
        ws = wb.create_sheet(name)
        ws.append(headers)
        for row in rows:
            for index, part in enumerate(cls._text_parts(row[-1])):
                values = [*row[:-1], part]
                if index and len(values) > 2:
                    values[-2] = f"{values[-2]} (continued)"
                ws.append(values)
        for cells in ws:
            for cell in cells:
                if cell.row > 1 and headers[cell.column - 1] == "Date" and isinstance(cell.value, str):
                    cell.value = date.fromisoformat(cell.value)
                    cell.number_format = "yyyy-mm-dd"
                if isinstance(cell.value, str):
                    cell.data_type = "s"
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                cell.font = Font(name="Calibri", size=11, color="1F2937")
                if cell.row > 1 and cell.column in money_columns and isinstance(cell.value, (Real, Decimal)):
                    cell.number_format = cls.MONEY_FORMAT
        for cell in ws[1]:
            cell.fill = PatternFill("solid", fgColor="23354D")
            cell.font = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
        for column, width in widths.items():
            ws.column_dimensions[column].width = width
        for cells in ws:
            lines = max(sum(max(1, ceil(len(line) / (widths.get(cell.column_letter, 20) - 3)))
                            for line in str(cell.value or "").split("\n")) for cell in cells)
            ws.row_dimensions[cells[0].row].height = max(24, min(409, lines * 15 + 6))
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        ws.sheet_properties.pageSetUpPr.fitToPage = True
        ws.page_setup.orientation = "portrait"
        ws.page_setup.paperSize = ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth = 1
        ws.page_setup.fitToHeight = 0
        ws.print_title_rows = "1:1"
        return ws

    def export_excel(self, reports, soft_entries=None):
        wb = Workbook()
        wb.remove(wb.active)
        summary = [["Period (Asia/Manila)", self._period_label(reports, soft_entries or [])]]
        summary.extend([label, sum(r.get(key, 0) for r in reports)]
                       for group, fields in self.REPORT_GROUPS.items() if group != "Activity" for label, key in fields)
        summary.extend([label, sum(r.get(key, 0) for r in reports)] for label, key in self.REPORT_GROUPS["Activity"])
        activity_start = len(summary) - len(self.REPORT_GROUPS["Activity"]) + 2
        summary.extend([f"{m.upper()} checkout sales", sum(r.get(f"{m}_total", 0) for r in reports)]
                       for m in self.METHODS if m != "unclassified" or any(r.get("unclassified_total", 0) for r in reports))
        self._sheet(wb, "Summary", ["Item", "Value"], summary, {"A": 43, "B": 28}, (2,))
        # Activity counts are not currency.
        for row in range(activity_start, activity_start + len(self.REPORT_GROUPS["Activity"])):
            wb["Summary"].cell(row, 2).number_format = "0"
        daily = [[r["report_date"], r["total_revenue"], r.get("total_collections", 0),
                  r.get("total_business_expenses", 0), r["net_balance"]] for r in reports]
        self._sheet(wb, "Daily Reports", ["Date", "Checkout Sales", "Debt Collected", "Business Expenses", "Net Movement"],
                    daily, {"A": 14, "B": 17, "C": 17, "D": 18, "E": 18}, (2, 3, 4, 5))
        details = [[r["report_date"], section, label, value]
                   for r in reports for section, items in self.report_sections(r).items() for label, value in items]
        ws = self._sheet(wb, "Report Details", ["Date", "Section", "Item", "Value"], details,
                         {"A": 14, "B": 23, "C": 35, "D": 28}, (4,))
        for row in ws.iter_rows(min_row=2):
            if row[1].value in {"Activity", "Checkout counts"}:
                row[3].number_format = "0"
        methods = [[r["report_date"], m["method"], label, m[flow]]
                   for r in reports for m in self._method_rows([r]) for flow, label in self.FLOW_LABELS.items()]
        self._sheet(wb, "Payment Methods", ["Date", "Method", "Flow", "Amount"], methods,
                    {"A": 14, "B": 16, "C": 43, "D": 20}, (4,))
        soft = [[s["balance_date"], f"{s['period']} #{s['id']}", f"{section}: {label}", value]
                for s in soft_entries or [] for section, items in self.soft_sections(s).items() for label, value in items]
        self._sheet(wb, "Soft Balances", ["Date", "Snapshot", "Item", "Value"], soft,
                    {"A": 14, "B": 15, "C": 40, "D": 28}, (4,))
        if len(reports) >= 2:
            current, previous = reports[:2]
            comparison = [[label, current.get(key, 0), previous.get(key, 0), current.get(key, 0) - previous.get(key, 0)]
                          for label, key in self.REPORT_GROUPS["Totals"] + self.REPORT_GROUPS["Sales"]]
            self._sheet(wb, "Period Comparison", ["Item", current["report_date"], previous["report_date"], "Change"],
                        comparison, {"A": 35, "B": 18, "C": 18, "D": 18}, (2, 3, 4))
        trends = [["Days represented", len(reports)]]
        if reports:
            trends += [["Average daily checkout sales", sum(r["total_revenue"] for r in reports) / len(reports)],
                       ["Average daily expenses", sum(r["total_expenses"] for r in reports) / len(reports)],
                       ["Highest checkout sales day", max(reports, key=lambda r: r["total_revenue"])["report_date"]]]
        self._sheet(wb, "Trends", ["Item", "Value"], trends, {"A": 40, "B": 28}, (2,))
        wb["Trends"]["B2"].number_format = "0"
        output = BytesIO()
        wb.save(output)
        output.seek(0)
        return send_file(output, mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                         as_attachment=True, download_name=f"daily_balance_{manila_date(datetime.utcnow())}.xlsx")
