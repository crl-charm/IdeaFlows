"""Already-paid missed checkouts must reconcile once, on the actual payment day."""
from datetime import date, datetime, timedelta
from csv import DictReader
from decimal import Decimal
from io import BytesIO, StringIO
from time import time
from types import SimpleNamespace
from uuid import uuid4

import pytest
from sqlalchemy import inspect, text
from openpyxl import load_workbook

from app import create_app, db, socketio
from app.core import get_notifier
from app.db.migrator import SchemaMigrator
from app.models import (CashBoxEvent, CheckoutVoidRequest, CustomerSession, FinanceBudget, FinanceTransaction,
    MenuItem, MissedCheckoutCorrection, Order, OrderItem, Receivable, SessionTimeEvent,
    SpaceType, StaffShift, Transaction, User)
from app.repositories.analytics_repository import AnalyticsRepository
from app.repositories.sales_repository import SalesRepository
from app.repositories.session_repository import SessionRepository
from app.services.cash_box_service import CashBoxService
from app.services.daily_balance_export_service import DailyBalanceExportService
from app.services.session_service import SessionService
from app.services.sales_service import SalesService
from app.services.analytics_oop_service import AnalyticsReportService
from app.utils.dates import MANILA_OFFSET


class Clock:
    value = datetime(2026, 10, 29, 3, 31)

    def now(self):
        return self.value


@pytest.fixture
def app():
    application = create_app()
    application.config.update(TESTING=True, WTF_CSRF_ENABLED=False, SINGLE_SESSION_ENABLED=False)
    return application


@pytest.fixture
def setup(app, monkeypatch):
    import app.routes.session_routes as routes
    clock = Clock()
    monkeypatch.setattr(routes, "_service", SessionService(SessionRepository(), clock, get_notifier()))
    with app.app_context():
        staff = db.session.get(User, 1)
        staff.full_name = "Diana"
        db.session.add(User(id=2, username="admin_gilbert", full_name="Admin Gilbert", role="admin", password=staff.password))
        premium = SpaceType(name="Premium Lounge", rate_per_minute=1, capacity=30)
        regular = SpaceType(name="Regular Lounge", rate_per_minute=1, capacity=30)
        boardroom = SpaceType(name="Boardroom", rate_per_minute=1, capacity=30)
        meal = MenuItem(name="Synthetic meal", price=128, inventory_mode="untracked")
        db.session.add_all([premium, regular, boardroom, meal])
        db.session.flush()
        row = CustomerSession(customer_name="Synthetic Rio", space_type_id=premium.id,
            time_in=datetime(2026, 10, 27, 9, 11), status="active",
            paused_at=datetime(2026, 10, 28, 3, 27, 26), paused_microseconds=0)
        db.session.add(row)
        db.session.flush()
        SessionService._record_time_event(row, "paused", row.paused_at, 1, "Diana", "staff",
            18 * 3600_000000 + 16 * 60_000000 + 26_000000, Decimal(370))
        order = Order(customer_session_id=row.id, handled_by=1, created_at=datetime(2026, 10, 27, 9, 20), status="done")
        db.session.add(order)
        db.session.flush()
        db.session.add(OrderItem(order_id=order.id, menu_item_id=meal.id, quantity=1, price=128,
                                name_snapshot=meal.name, status="done"))
        db.session.commit()
        session_id = row.id
    admin = app.test_client()
    with admin.session_transaction() as auth:
        auth.update(user_id=2, username="admin_gilbert", role="admin", last_activity=time())
    return admin, clock, session_id


def details(**changes):
    result = dict(departure_at="2026-10-27T19:11", paid_at="2026-10-27T19:11",
                  amount_received="168.00", payment_method="cash", collector_id=1, reason="outage", note="")
    result.update(changes)
    return result


def preview(admin, session_id, data=None):
    return admin.post(f"/api/sessions/{session_id}/missed-checkout/preview", json=data or details())


def complete(admin, session_id, data=None, token=None, key=None):
    data = data or details()
    if token is None:
        response = preview(admin, session_id, data)
        assert response.status_code == 200, response.get_json()
        token = response.get_json()["preview_token"]
    return admin.post(f"/api/sessions/{session_id}/missed-checkout", json={
        **data, "preview_token": token, "payment_confirmed": True,
    }, headers={"Idempotency-Key": key or str(uuid4())})


def test_rio_shaped_correction_ignores_next_day_pause_and_dates_one_payment(app, setup):
    admin, clock, session_id = setup
    quote = preview(admin, session_id)
    assert quote.status_code == 200, quote.get_json()
    data = quote.get_json()
    assert data["before"]["total"] == "498.00"
    assert data["total"] == 168 and data["time_bill"] == 40 and data["food_bill"] == 128
    assert data["billable_seconds"] == 7200 and data["matches"] is True
    with app.app_context():
        assert Transaction.query.count() == MissedCheckoutCorrection.query.count() == 0
    assert admin.post(f"/api/checkout/{session_id}", json={"payment_method":"cash","amount_tendered":"168"}).status_code == 400
    clock.value += timedelta(seconds=15)  # Reviewing the quote does not itself stale it.
    saved = complete(admin, session_id, token=data["preview_token"])
    assert saved.status_code == 200, saved.get_json()
    assert saved.get_json()["total_bill"] == 168
    with app.app_context():
        tx, audit = Transaction.query.one(), MissedCheckoutCorrection.query.one()
        row = db.session.get(CustomerSession, session_id)
        assert tx.paid_at == datetime(2026, 10, 27, 11, 11)
        assert tx.created_at == clock.now() and tx.billing_end_at == tx.paid_at
        assert tx.amount_tendered is None and tx.change_given is None
        assert tx.collected_by == "Diana" and audit.collector_id == 1
        assert audit.actor_id == 2 and audit.actor_name == "admin_gilbert"
        assert row.status == "completed" and row.paused_at is None and row.paused_microseconds == 0
        assert OrderItem.query.one().price == 128 and OrderItem.query.one().quantity == 1
        ledger = SalesRepository().daily_ledger(date(2026,10,27), date(2026,10,29))
        assert ledger[date(2026,10,27)]["total_revenue"] == 168
        assert ledger[date(2026,10,27)]["cowork_bill"] == 40
        assert ledger[date(2026,10,27)]["cafe_bill"] == 128
        assert date(2026,10,29) not in ledger
        assert ledger[date(2026,10,27)]["total_refunds"] == 0
        assert len(ledger[date(2026,10,27)]["late_checkouts"]) == 1
        assert SalesRepository().summary_for_day(date(2026,10,27)).total_revenue == 168
        assert SalesRepository().summary_for_day(date(2026,10,29)).total_revenue == 0
    record = admin.get("/api/checkout-records?paginated=1&date_from=2026-10-27&date_to=2026-10-27").get_json()["data"][0]
    assert record["payment_date"] == "2026-10-27" and record["created_date"] == "2026-10-29"
    assert record["recorded_late"] is True and record["seconds_spent"] == 7200
    assert admin.get("/api/checkout-records?paginated=1&date_from=2026-10-29&date_to=2026-10-29").get_json()["data"] == []
    receipt = admin.get(saved.get_json()["receipt_url"])
    assert receipt.status_code == 200 and b"payment already received" in receipt.data
    history = admin.get(f"/api/sessions/{session_id}/time-history").get_json()
    assert history["events"][0]["action"] == "paused" and history["events"][1]["action"] == "missed_checkout"
    assert history["correction"]["received_by"] == "Diana"
    assert "2026-10-29" in history["correction"]["recorded_at"]
    assert app.test_client().get("/api/current-customers").get_json()["customers"] == []


@pytest.mark.parametrize("change", [
    {"departure_at": "2026-10-27T17:00"}, {"departure_at": "2026-10-30T09:00"},
    {"departure_at": []}, {"paid_at": "2026-10-27T17:00"}, {"paid_at": "2026-10-30T10:00"},
    {"amount_received": "nan"}, {"amount_received": "-1"}, {"payment_method": []},
    {"collector_id": True}, {"collector_id": 999}, {"reason": "other", "note": ""},
    {"reason": []}, {"discount_type": []},
])
def test_invalid_correction_cannot_mutate_the_visit(app, setup, change):
    admin, _, session_id = setup
    assert preview(admin, session_id, details(**change)).status_code == 400
    with app.app_context():
        assert Transaction.query.count() == MissedCheckoutCorrection.query.count() == 0
        assert db.session.get(CustomerSession, session_id).status == "active"


def test_mismatch_requires_reconciliation_not_a_discount_or_debt(app, setup):
    admin, _, session_id = setup
    quote = preview(admin, session_id, details(amount_received="160")).get_json()
    assert quote["matches"] is False and quote["difference"] == -8
    result = complete(admin, session_id, details(amount_received="160"), quote["preview_token"])
    assert result.status_code == 400
    with app.app_context():
        assert Transaction.query.count() == Receivable.query.count() == MissedCheckoutCorrection.query.count() == 0


def test_admin_only_and_preview_revision_prevents_changed_orders(app, setup):
    admin, _, session_id = setup
    anonymous = app.test_client()
    assert preview(anonymous, session_id).status_code == 401
    with anonymous.session_transaction() as auth:
        auth.update(user_id=1, username="test_user", role="staff", last_activity=time())
    assert preview(anonymous, session_id).status_code == 403
    assert complete(anonymous, session_id, token="forged").status_code == 403
    quote = preview(admin, session_id).get_json()
    assert complete(admin, session_id, token="forged").status_code == 409
    with app.app_context():
        OrderItem.query.one().quantity = 2
        db.session.commit()
    result = complete(admin, session_id, token=quote["preview_token"])
    assert result.status_code == 409 and "changed" in result.get_json()["error"]
    with app.app_context():
        assert Transaction.query.count() == 0


@pytest.mark.parametrize("pairs,open_pause,expected", [
    ([], None, 7200),
    ([("2026-10-27T17:41", "2026-10-27T18:11")], None, 5400),
    ([("2026-10-27T18:41", "2026-10-27T19:41")], None, 5400),
    ([("2026-10-28T09:00", "2026-10-28T10:00")], None, 7200),
    ([("2026-10-27T17:41", "2026-10-27T18:11")], "2026-10-28T09:00", 5400),
])
def test_historical_pause_intervals(app, setup, pairs, open_pause, expected):
    from app.utils.dates import entered_manila_time
    admin, _, session_id = setup
    with app.app_context():
        SessionTimeEvent.query.delete()
        row = db.session.get(CustomerSession, session_id)
        row.paused_at = entered_manila_time(open_pause) if open_pause else None
        row.paused_microseconds = 0
        for start, end in pairs:
            start, end = entered_manila_time(start), entered_manila_time(end)
            for action, when in (("paused",start),("resumed",end)):
                db.session.add(SessionTimeEvent(session_id=session_id,action=action,occurred_at=when,
                    business_date=(when+MANILA_OFFSET).date(),actor_id=1,actor_name="Diana",actor_role="staff",
                    billable_microseconds_before=0,billable_microseconds_after=0,time_bill_before=0,time_bill_after=0))
            row.paused_microseconds += (end-start)//timedelta(microseconds=1)
        if open_pause:
            db.session.add(SessionTimeEvent(session_id=session_id,action="paused",occurred_at=row.paused_at,
                business_date=(row.paused_at+MANILA_OFFSET).date(),actor_id=1,actor_name="Diana",actor_role="staff",
                billable_microseconds_before=0,billable_microseconds_after=0,time_bill_before=0,time_bill_after=0))
        db.session.commit()
    result = preview(admin, session_id)
    assert result.status_code == 200, result.get_json()
    assert result.get_json()["billable_seconds"] == expected


def test_missing_pause_history_and_food_after_departure_require_review(app, setup):
    admin, _, session_id = setup
    with app.app_context():
        SessionTimeEvent.query.delete()
        db.session.commit()
    assert preview(admin, session_id).status_code == 400
    with app.app_context():
        row = db.session.get(CustomerSession,session_id)
        row.paused_at = None
        Order.query.one().created_at = datetime(2026,10,28,1)
        db.session.commit()
    result = preview(admin,session_id)
    assert result.status_code == 400 and "Food was ordered after" in result.get_json()["error"]


def test_atomic_retry_and_notification_failure(app, setup, monkeypatch):
    admin, _, session_id = setup
    token = preview(admin,session_id).get_json()["preview_token"]
    original = SessionRepository.commit
    def failed_commit(self):
        db.session.flush()
        raise RuntimeError("Simulated save failure")
    monkeypatch.setattr(SessionRepository,"commit",failed_commit)
    key = str(uuid4())
    assert complete(admin,session_id,token=token,key=key).status_code == 500
    with app.app_context():
        assert Transaction.query.count() == MissedCheckoutCorrection.query.count() == 0
        assert db.session.get(CustomerSession,session_id).status == "active"
    monkeypatch.setattr(SessionRepository,"commit",original)
    def failed_emit(*args,**kwargs):
        raise RuntimeError("Simulated socket outage")
    monkeypatch.setattr(socketio,"emit",failed_emit)
    assert complete(admin,session_id,token=token,key=key).status_code == 200
    assert complete(admin,session_id,token=token,key=key).status_code == 409
    assert complete(admin,session_id,token=token).status_code == 409
    with app.app_context():
        assert Transaction.query.count() == MissedCheckoutCorrection.query.count() == 1


def test_digital_payment_day_can_differ_from_departure(app, setup):
    admin, _, session_id = setup
    result = complete(admin,session_id,details(payment_method="gcash",paid_at="2026-10-28T09:00"))
    assert result.status_code == 200
    with app.app_context():
        ledger = SalesRepository().daily_ledger(date(2026,10,27),date(2026,10,29))
        assert ledger.get(date(2026,10,27), {}).get("total_revenue", 0) == 0
        assert ledger[date(2026,10,28)]["gcash_total"] == 168
        assert ledger[date(2026,10,28)]["cash_box_movements"]["customer"] == 0


def test_known_cash_tender_keeps_net_receipt_and_change_separate(app, setup):
    admin, _, session_id = setup
    assert complete(admin, session_id, details(amount_tendered="200")).status_code == 200
    with app.app_context():
        tx = Transaction.query.one()
        assert tx.total_bill == 168 and tx.amount_tendered == 200 and tx.change_given == 32
        assert SalesRepository().daily_ledger(date(2026,10,27),date(2026,10,27))[date(2026,10,27)]["cash_box_receipts"]["customer"] == 168


@pytest.mark.parametrize("opening_day,opening_cash,closing", [
    (date(2026,10,26), 0, 168),
    (date(2026,10,28), 168, 168),  # That physical opening count already contains the payment.
])
def test_backdated_cash_not_added_twice_after_a_physical_opening(app, setup, opening_day, opening_cash, closing):
    admin, _, session_id = setup
    with app.app_context():
        for box, amount in (("petty",300),("customer",opening_cash)):
            db.session.add(CashBoxEvent(kind="opening",box=box,amount=amount,effect=amount,
                effective_date=opening_day,actor_id=2,actor_name="Admin Gilbert"))
        db.session.commit()
    assert complete(admin,session_id).status_code == 200
    with app.app_context():
        boxes = CashBoxService.summaries(date(2026,10,29))
        assert boxes[date(2026,10,29)]["closing"] == {"petty":300,"customer":closing}
        assert CashBoxEvent.query.count() == 2  # No extra top-up for this receipt.


def test_later_refund_uses_actual_payment_day_and_preserves_correction(app, setup):
    admin, _, session_id = setup
    assert complete(admin,session_id).status_code == 200
    with app.app_context():
        tx_id = Transaction.query.one().id
    requested = admin.post(f"/api/checkout-records/{tx_id}/void-request",json={"reason":"Refund agreed"})
    assert requested.status_code == 201
    request_id = requested.get_json()["request_id"]
    approved = admin.post(f"/api/checkout-records/void-requests/{request_id}/approve",json={
        "money_treatment":"refund","decision_reason":"Money returned to customer"})
    assert approved.status_code == 200, approved.get_json()
    with app.app_context():
        assert db.session.get(CheckoutVoidRequest,request_id).original_business_date == date(2026,10,27)
        ledger = SalesRepository().daily_ledger(date(2026,10,27),date(2026,10,29))
        assert ledger[date(2026,10,27)]["total_revenue"] == 168
        assert ledger[date(2026,10,29)]["total_revenue"] == 0
        assert ledger[date(2026,10,29)]["total_refunds"] == 168
        assert MissedCheckoutCorrection.query.count() == 1


def test_reports_analytics_and_exports_use_payment_day_with_late_entry_audit(app, setup):
    admin, _, session_id = setup
    assert complete(admin,session_id).status_code == 200
    with app.app_context():
        repo = SalesRepository()
        report = SalesService(repo).list_reports(date(2026,10,27),date(2026,10,27))[0]
        assert report["report_date"] == "2026-10-27" and report["total_revenue"] == 168
        assert repo.get_daily_transactions(date(2026,10,29)) == []
        assert len(repo.get_daily_transactions(date(2026,10,27))) == 1
        assert AnalyticsReportService(db).get_summary("2026-10-27","2026-10-27")["daily_sales"] == [
            {"date":"2026-10-27","revenue":168,"transactions":1}]
        assert AnalyticsReportService(db).get_summary("2026-10-29","2026-10-29")["daily_sales"] == []
        assert any(item == {"date":"2026-10-27","revenue":168} for item in AnalyticsRepository().daily_revenue())
        assert AnalyticsRepository().peak_hours() == [{"hour":19,"count":1}]
        service = DailyBalanceExportService(db)
        context = service.build_pdf_context([report],[])
        assert context["total_revenue"] == 168
        assert "2026-10-29" in context["report_details"][0][1]["Late-recorded payments"][0][1]
        with app.test_request_context():
            response = service.export_csv([report],[])
            response.direct_passthrough = False
            rows = list(DictReader(StringIO(response.get_data().decode("utf-8-sig"))))
            late = next(row for row in rows if row["Section"].endswith("Late-recorded payments"))
            assert late["Date"] == "2026-10-27" and "recorded 2026-10-29" in late["Notes"]
            response = service.export_excel([report],[])
            response.direct_passthrough = False
            wb = load_workbook(BytesIO(response.get_data()))
            assert wb["Daily Reports"]["A2"].value.date() == date(2026,10,27)
            assert wb["Daily Reports"]["B2"].value == 168
            assert any(row[1].value == "Late-recorded payments" for row in wb["Report Details"].iter_rows(min_row=2))


@pytest.mark.parametrize("unsupported", ["boardroom","food_only","credit","debt"])
def test_conflicting_financial_states_require_existing_correction_flow(app, setup, unsupported):
    admin, _, session_id = setup
    with app.app_context():
        row = db.session.get(CustomerSession,session_id)
        if unsupported == "boardroom":
            row.space_type_id = SpaceType.query.filter_by(name="Boardroom").one().id
        elif unsupported == "food_only":
            row.service_mode = "food_only"
        elif unsupported == "credit":
            row.credit_balance = 20
        else:
            db.session.add(Receivable(session_id=session_id,customer_name=row.customer_name,
                items_description="Already recorded debt",amount_owed=168,due_date=date(2026,10,27),created_by=1))
        db.session.commit()
    assert preview(admin,session_id).status_code in (400,409)
    with app.app_context():
        assert Transaction.query.count() == MissedCheckoutCorrection.query.count() == 0


def test_supported_discount_uses_saved_food_price_and_cannot_override_actor(app, setup):
    admin, _, session_id = setup
    with app.app_context():
        item_id = OrderItem.query.one().id
    data = details(discount_type="pwd",discount_item_id=item_id,amount_received="142.40",actor_id=1,actor_role="staff")
    result = complete(admin,session_id,data)
    assert result.status_code == 200, result.get_json()
    with app.app_context():
        tx, audit = Transaction.query.one(), MissedCheckoutCorrection.query.one()
        assert tx.food_bill == 128 and tx.discount_amount == Decimal("25.60") and tx.total_bill == Decimal("142.40")
        assert audit.actor_id == 2 and audit.actor_role == "admin"


def test_canonical_admin_and_past_staff_collectors_remain_distinct(app, setup):
    admin, _, session_id = setup
    with app.app_context():
        db.session.get(User,2).is_active = False  # Canonical shadow of an admin account.
        db.session.get(User,1).is_active = False  # Staff who later left the business.
        db.session.commit()
    options = admin.get("/api/missed-checkout/options").get_json()["collectors"]
    assert next(item for item in options if item["id"]==2)["active"] is True
    assert next(item for item in options if item["id"]==1)["active"] is False
    assert preview(admin,session_id).status_code == 200  # Historical collector is valid.


def test_new_possible_payment_after_preview_requires_another_review(app, setup):
    admin, _, session_id = setup
    token = preview(admin,session_id).get_json()["preview_token"]
    with app.app_context():
        budget = FinanceBudget(_name="Synthetic budget")
        db.session.add(budget)
        db.session.flush()
        db.session.add(FinanceTransaction(_budget_id=budget.id,_type="income",_amount=168,_description="Possible prior receipt",
            created_at=datetime(2026,10,27,11,11)))
        db.session.commit()
    assert complete(admin,session_id,token=token).status_code == 409
    refreshed = preview(admin,session_id).get_json()
    assert any("Manual income" in item for item in refreshed["possible_existing_payments"])


def test_normal_checkout_winning_after_preview_blocks_duplicate_sale(app, setup):
    admin, _, session_id = setup
    token = preview(admin,session_id).get_json()["preview_token"]
    assert admin.post(f"/api/checkout/{session_id}",json={"payment_method":"cash","amount_tendered":"498"}).status_code == 200
    assert complete(admin,session_id,token=token).status_code == 409
    with app.app_context():
        assert Transaction.query.count() == 1 and MissedCheckoutCorrection.query.count() == 0


@pytest.mark.parametrize("arrival,paid_day,paid_time", [
    ("2026-09-30", "2026-09-30", "23:59"),
    ("2026-09-30", "2026-10-01", "00:01"),
    ("2025-12-31", "2026-01-01", "00:01"),
])
def test_prior_month_year_and_manila_midnight_use_actual_payment_date(app, setup, arrival, paid_day, paid_time):
    admin, _, session_id = setup
    shift = date.fromisoformat(arrival) - date(2026,10,27)
    with app.app_context():
        row = db.session.get(CustomerSession,session_id)
        row.time_in += shift
        row.paused_at += shift
        SessionTimeEvent.query.one().occurred_at += shift
        Order.query.one().created_at += shift
        db.session.commit()
    result = complete(admin,session_id,details(departure_at=f"{arrival}T19:11",paid_at=f"{paid_day}T{paid_time}"))
    assert result.status_code == 200, result.get_json()
    with app.app_context():
        day = date.fromisoformat(paid_day)
        assert SalesRepository().daily_ledger(day,day)[day]["total_revenue"] == 168
        assert Transaction.query.one().created_at == datetime(2026,10,29,3,31)


def test_historical_pause_precision_uses_exact_aggregate_or_stops_at_ambiguous_boundary(app):
    from app.utils.billing import historical_billable_microseconds
    start = datetime(2026,10,27,9)
    departure = start + timedelta(minutes=70)
    row = SimpleNamespace(time_in=start, paused_at=None, paused_microseconds=600_500000,
                          space_type=SimpleNamespace(name="Premium Lounge"))
    events = [SimpleNamespace(action="paused",occurred_at=start+timedelta(minutes=10)),
              SimpleNamespace(action="resumed",occurred_at=start+timedelta(minutes=20))]
    assert historical_billable_microseconds(row,departure,events) == 3599_500000
    # A completed pause after departure makes the old fractional allocation
    # uncertain; do not guess when it could change the started billing block.
    events += [SimpleNamespace(action="paused",occurred_at=start+timedelta(minutes=80)),
               SimpleNamespace(action="resumed",occurred_at=start+timedelta(minutes=90))]
    row.paused_microseconds += 600_000000
    with pytest.raises(ValueError,match="precision affects"):
        historical_billable_microseconds(row,departure,events)


def test_late_entry_preserves_saved_daily_and_soft_cash_snapshots(app, setup):
    from app.models import DailySalesReport, SoftBalanceEntry
    admin, _, session_id = setup
    with app.app_context():
        for box in ("petty","customer"):
            db.session.add(CashBoxEvent(kind="opening",box=box,amount=0,effect=0,
                effective_date=date(2026,10,26),actor_id=2,actor_name="Admin Gilbert"))
        db.session.commit()
        service = SalesService(SalesRepository())
        service.generate_report(date(2026,10,27),generated_by=2,notes="Original closing count")
        service.create_soft_balance(date(2026,10,27),period="PM",generated_by=2,notes="Original PM count")
    assert complete(admin,session_id).status_code == 200
    with app.app_context():
        assert DailySalesReport.query.one().total_revenue == SoftBalanceEntry.query.one().total_revenue == 0
        service = SalesService(SalesRepository())
        report = service.list_reports(date(2026,10,27),date(2026,10,27))[0]
        assert report["total_revenue"] == 168
        assert report["snapshot_customer_cash_closing"] == 0
        assert report["customer_cash_snapshot_difference"] == 168
        soft = service.list_soft_balances()[0]
        assert soft["customer_cash_closing"] == 0 and soft["customer_cash_live_difference"] == 168


def test_handover_reminder_rechecks_active_customers_without_changing_bills(app, setup):
    _, _, session_id = setup
    staff = app.test_client()
    with staff.session_transaction() as auth:
        auth.update(user_id=1,username="test_user",role="staff",last_activity=time())
    with app.app_context():
        shift = StaffShift(user_id=1,shift_role="cashier",time_in=datetime.utcnow()-timedelta(hours=2))
        db.session.add(shift)
        db.session.commit()
        shift_id = shift.id
    payload = {"time_out":(datetime.utcnow()+MANILA_OFFSET-timedelta(minutes=1)).strftime("%Y-%m-%dT%H:%M")}
    warning = staff.post("/api/staff-shift/time-out",json=payload)
    assert warning.status_code == 409 and warning.get_json()["needs_handover"] is True
    with app.app_context():
        assert db.session.get(StaffShift,shift_id).time_out is None
        row = db.session.get(CustomerSession,session_id)
        old_pause = row.paused_at
        db.session.add(CustomerSession(customer_name="New arrival",space_type_id=row.space_type_id,
                       time_in=datetime.utcnow(),status="active"))
        db.session.commit()
    payload["handover_token"] = warning.get_json()["handover_token"]
    changed = staff.post("/api/staff-shift/time-out",json=payload)
    assert changed.status_code == 409 and len(changed.get_json()["active_customers"]) == 2
    payload["handover_token"] = changed.get_json()["handover_token"]
    assert staff.post("/api/staff-shift/time-out",json=payload).status_code == 200
    with app.app_context():
        row = db.session.get(CustomerSession,session_id)
        assert row.status == "active" and row.paused_at == old_pause
        assert Transaction.query.count() == 0


def test_migration_is_additive_and_legacy_payment_dates_stay_unchanged(app, setup):
    _, _, session_id = setup
    with app.app_context():
        tx = Transaction(session_id=session_id,time_bill=20,food_bill=0,total_bill=20,
                         created_at=datetime(2026,9,30,16,5))
        db.session.add(tx)
        db.session.commit()
        tx_id = tx.id
        db.session.execute(text("DROP INDEX IF EXISTS idx_transactions_paid_at"))
        db.session.execute(text("ALTER TABLE transactions DROP COLUMN paid_at"))
        db.session.execute(text("DROP TABLE missed_checkout_corrections"))
        db.session.commit()
        SchemaMigrator(db,app).run()
        SchemaMigrator(db,app).run()
        db.session.expire_all()
        tx = db.session.get(Transaction,tx_id)
        assert tx.paid_at is None and tx.effective_paid_at == datetime(2026,9,30,16,5)
        assert inspect(db.engine).has_table("missed_checkout_corrections")
        payment_index = next(item for item in inspect(db.engine).get_indexes("transactions") if item["name"] == "idx_transactions_paid_at")
        assert payment_index["column_names"] == ["paid_at","created_at"]
        assert SalesRepository().daily_ledger(date(2026,10,1),date(2026,10,1))[date(2026,10,1)]["total_revenue"] == 20
