"""The live screen, snapshot and export must agree on one Manila business day."""

from datetime import date, datetime, timezone
from decimal import Decimal
from io import BytesIO

import pytest
from flask import render_template
from openpyxl import load_workbook
from sqlalchemy import create_engine, inspect, text

from app import create_app, db
from app.models import (
    CheckoutVoidRequest, CustomerSession, Expense, FinanceBudget, FinanceTransaction, Order,
    Payable, PayablePayment, Receivable, ReceivablePayment, SpaceType, Transaction,
)
from app.repositories.sales_repository import SalesRepository
from app.services.daily_balance_export_service import DailyBalanceExportService
from app.services.sales_service import SalesService
from app.services.finance_oop_service import FinanceService
from app.utils.dates import manila_date
from app.db.daily_balance_upgrade import upgrade


@pytest.fixture
def app():
    application = create_app()
    application.config.update(TESTING=True, WTF_CSRF_ENABLED=False)
    return application


def test_pdf_hides_empty_unknown_method_without_losing_older_amounts():
    assert all(row["method"] != "unclassified" for row in DailyBalanceExportService._method_rows([{}]))
    rows = DailyBalanceExportService._method_rows([{"expense_methods": {"unclassified": 5}}])
    assert next(row for row in rows if row["method"] == "unclassified")["expense"] == 5


def test_live_ledger_reconciles_midnight_payments_and_export(app):
    first_day = date(2026, 9, 23)
    business_day = date(2026, 9, 24)
    with app.app_context():
        space = SpaceType(name="Regular Lounge", rate_per_minute=Decimal("0.1667"))
        db.session.add(space)
        db.session.flush()
        sessions = [CustomerSession(customer_name=f"Customer {i}", space_type_id=space.id,
                                    time_in=datetime(2026, 9, 23, 15, 0), status="completed") for i in (1, 2)]
        db.session.add_all(sessions)
        db.session.flush()
        db.session.add_all([
            Transaction(session_id=sessions[0].id, time_bill=Decimal("10"), food_bill=0,
                        total_bill=Decimal("10"), payment_method="cash",
                        created_at=datetime(2026, 9, 23, 15, 59)),
            Transaction(session_id=sessions[1].id, time_bill=Decimal("50"), food_bill=0,
                        total_bill=Decimal("50"), payment_method="gcash",
                        created_at=datetime(2026, 9, 23, 16, 10)),
        ])
        debts = [Receivable(customer_name="Debtor", items_description=f"Debt {i}",
                            amount_owed=Decimal("50"), partial_paid=Decimal(str(amount)),
                            due_date=business_day, created_by=1)
                 for i, amount in ((1, 30), (2, 20))]
        db.session.add_all(debts)
        db.session.flush()
        db.session.add_all([
            ReceivablePayment(receivable_id=debts[0].id, amount=Decimal("30"), payment_method="cash",
                              received_by=1, payment_group_id="one-payment",
                              received_at=datetime(2026, 9, 23, 16, 15)),
            ReceivablePayment(receivable_id=debts[1].id, amount=Decimal("20"), payment_method="cash",
                              received_by=1, payment_group_id="one-payment",
                              received_at=datetime(2026, 9, 23, 16, 15)),
            Expense(category="supplies", description="Old method unknown", amount=Decimal("5"),
                    expense_date=business_day, logged_by=1),
            Payable(creditor_name="Supplier", items_description="Stock", amount_owed=Decimal("50"),
                    due_date=business_day, incurred_date=business_day, created_by=1),
        ])
        db.session.commit()
        payable_id = Payable.query.one().id

    client = app.test_client()
    with client.session_transaction() as auth:
        auth.update(user_id=1, username="test_user", role="admin",
                    last_activity=datetime.now(timezone.utc).timestamp())

    expense = client.post("/admin/expenses/api/expenses", json={
        "category": "supplies", "description": "Paper", "amount": "10.00",
        "expense_date": business_day.isoformat(), "payment_method": "cash",
    }, headers={"Idempotency-Key": "ledger-expense"})
    assert expense.status_code == 201
    expense_id = expense.get_json()["data"]["id"]
    with app.app_context():
        db.session.get(Expense, expense_id).created_at = datetime(2026, 9, 23, 16, 17)
        db.session.commit()
    payment_body = {"amount": "20.00", "payment_method": "cash"}
    payment_url = f"/admin/payables/api/payables/{payable_id}/mark-paid"
    paid = client.patch(payment_url, json=payment_body, headers={"Idempotency-Key": "ledger-payable"})
    retry = client.patch(payment_url, json=payment_body, headers={"Idempotency-Key": "ledger-payable"})
    assert paid.status_code == 200
    assert retry.status_code == 200
    assert retry.get_json()["payment_id"] == paid.get_json()["payment_id"]
    assert retry.get_json()["replayed"] is True
    with app.app_context():
        payment = PayablePayment.query.one()
        payment.paid_at = datetime(2026, 9, 23, 16, 20)
        db.session.commit()

    response = client.get("/admin/daily-balance/api/reports?start_date=2026-09-23&end_date=2026-09-24")
    assert response.status_code == 200
    rows = {r["report_date"]: r for r in response.get_json()["data"]}
    assert rows[first_day.isoformat()]["cash_total"] == 10
    day = rows[business_day.isoformat()]
    assert day["total_revenue"] == 50
    assert day["gcash_total"] == 50
    assert day["cash_total"] == 0
    assert day["total_collections"] == 50
    assert day["collection_count"] == 1
    assert day["total_expenses"] == 15
    assert day["total_business_expenses"] == 15
    assert day["total_external_expenses"] == 0
    assert day["expense_methods"]["unclassified"] == 5
    assert day["total_payables_paid"] == 20
    assert day["cash_on_hand"] == 20
    assert day["net_balance"] == 65
    assert day["generated_by"] == "Not generated"

    with app.app_context():
        payment = PayablePayment.query.one()
        assert (payment.amount, payment.payment_method, payment.paid_by,
                payment.balance_before, payment.balance_after) == (
                    Decimal("20"), "cash", 1, Decimal("50"), Decimal("30"))
        assert SalesRepository().summary_for_day(business_day).total_revenue == 50
        SalesService(SalesRepository()).generate_report(business_day, 1, "Audit snapshot")
        assert SalesService(SalesRepository()).list_reports(business_day, business_day)[0]["net_balance"] == 65
        context = DailyBalanceExportService(db).build_pdf_context([day], [])
        assert context["total_collections"] == 50
        assert context["method_rows"][0]["net"] == 20
        assert next(row for row in context["method_rows"] if row["method"] == "unclassified")["expense"] == 5

    export = client.get("/admin/daily-balance/api/reports/export-csv?start_date=2026-09-24&end_date=2026-09-24")
    assert export.status_code == 200
    assert "Debt Collected" in export.data.decode("utf-8")
    assert "₱65.00" in export.data.decode("utf-8")

    voided = client.delete(f"/admin/expenses/api/expenses/{expense_id}", json={"reason": "Duplicate receipt"},
                           headers={"Idempotency-Key": "ledger-expense-void"})
    assert voided.status_code == 200
    with app.app_context():
        db.session.get(Expense, expense_id).voided_at = datetime(2026, 9, 23, 16, 30)
        db.session.commit()
    refreshed = client.get("/admin/daily-balance/api/reports?start_date=2026-09-24&end_date=2026-09-24")
    assert refreshed.get_json()["data"][0]["net_balance"] == 75
    history = client.get("/admin/expenses/api/expenses").get_json()["data"]
    assert next(e for e in history if e["id"] == expense_id)["void_reason"] == "Duplicate receipt"


def test_budget_adjustment_is_separate_from_checkout_sales(app):
    with app.app_context():
        db.session.add(FinanceBudget(_name="Operating", _total_budget=Decimal("1000")))
        db.session.commit()
        budget_id = FinanceBudget.query.one().id

    client = app.test_client()
    with client.session_transaction() as auth:
        auth.update(user_id=1, username="test_user", role="admin",
                    last_activity=datetime.now(timezone.utc).timestamp())
    for kind, amount, method, key in (
        ("income", "50.00", "bdo", "other-income"),
        ("expense", "20.00", "cash", "budget-spend"),
    ):
        response = client.post("/api/finance/transaction", json={
            "budget_id": budget_id, "type": kind, "amount": amount,
            "description": f"{kind} correction", "payment_method": method,
        }, headers={"Idempotency-Key": key})
        assert response.status_code == 200

    with app.app_context():
        assert FinanceTransaction.query.count() == 2
        day = manila_date(FinanceTransaction.query.first().created_at)
        rows = SalesService(SalesRepository()).list_reports(day, day)
        row = rows[0]
        assert row["total_revenue"] == 0
        assert row["total_other_income"] == 50
        assert row["total_budget_spend"] == 20
        assert row["net_balance"] == 30
        assert row["cash_on_hand"] == -20
        assert row["adjustment_income_methods"]["bdo"] == 50
        assert row["adjustment_expense_methods"]["cash"] == 20
        assert FinanceService(db).get_summary()["transactions"][0]["actor"] == "test_user"
        with pytest.raises(ValueError, match="already recorded"):
            FinanceService(db).add_transaction(budget_id, "income", "50.00", "Retry", "bdo", 1, "other-income")


def test_cowork_cafe_discount_split_reconciles_with_daily_balance_and_exports(app):
    first_day, next_day = date(2026, 10, 5), date(2026, 10, 6)
    before_midnight = datetime(2026, 10, 5, 15, 59)  # 23:59 in Manila
    after_midnight = datetime(2026, 10, 5, 16, 1)
    with app.app_context():
        regular = SpaceType(name="Regular Lounge", rate_per_minute=Decimal("0.1667"))
        takeout = SpaceType(name="Take Out", rate_per_minute=Decimal("0"))
        db.session.add_all([regular, takeout])
        db.session.flush()

        def checkout(name, time, food, discount, method, *, at=before_midnight,
                     food_only=False, food_discount=False, credit_applied=0):
            session = CustomerSession(customer_name=name, space_type_id=takeout.id if food_only else regular.id,
                                      service_mode="food_only" if food_only else "timed", status="completed")
            db.session.add(session)
            db.session.flush()
            tx = Transaction(session_id=session.id, time_bill=Decimal(str(time)), food_bill=Decimal(str(food)),
                             discount_amount=Decimal(str(discount)), discount_item_id=1 if food_discount else None,
                             total_bill=Decimal(str(time + food - discount)), payment_method=method,
                             credit_applied=Decimal(str(credit_applied)), created_at=at)
            db.session.add(tx)
            db.session.flush()
            return tx

        checkout("Mixed", 20, 180, 18, "cash", food_discount=True)
        checkout("Time discount", 50, 0, 10, "gcash")
        checkout("Take Out", 0, 100, 0, "bdo", food_only=True)
        checkout("Bank food", 0, 20, 0, "bpi", food_only=True)
        checkout("Bank time", 30, 0, 0, "queenbank")
        refunded = checkout("Refunded tomorrow", 0, 100, 0, "cash", food_only=True)
        mistaken = checkout("False entry", 50, 0, 0, "cash")
        credited = checkout("Credit retained", 0, 60, 0, "gcash", food_only=True)
        pending = checkout("Pending void", 0, 7, 0, "cash", food_only=True)
        rejected = checkout("Rejected void", 0, 8, 0, "gcash", food_only=True)
        checkout("Credit replacement", 0, 60, 0, "cash", at=after_midnight,
                 food_only=True, credit_applied=60)

        for tx, status, treatment, refund in (
            (refunded, "approved", "refund", 100),
            (mistaken, "approved", "false_entry", 0),
            (credited, "approved", "retained_credit", 0),
            (pending, "pending", None, 0),
            (rejected, "rejected", None, 0),
        ):
            db.session.add(CheckoutVoidRequest(
                transaction_id=tx.id, status=status, money_treatment=treatment,
                request_reason="Audit", requested_by_id=1, original_amount=tx.total_bill,
                original_payment_method=tx.payment_method, original_business_date=first_day,
                decided_at=after_midnight if status != "pending" else None,
                decision_business_date=next_day if status == "approved" else None,
                refund_amount=Decimal(str(refund)),
            ))

        unpaid = CustomerSession(customer_name="Unpaid order", space_type_id=takeout.id,
                                 service_mode="food_only", status="active")
        db.session.add(unpaid)
        db.session.flush()
        db.session.add(Order(customer_session_id=unpaid.id, food_total_after=Decimal("200"),
                             created_at=before_midnight))
        debt = Receivable(customer_name="Debtor", items_description="Old debt", amount_owed=Decimal("150"),
                          partial_paid=Decimal("150"), due_date=next_day, created_by=1)
        db.session.add(debt)
        db.session.flush()
        db.session.add(ReceivablePayment(receivable_id=debt.id, amount=Decimal("150"),
                                         payment_method="cash", received_by=1, received_at=after_midnight))
        db.session.commit()

        rows = SalesService(SalesRepository()).list_reports(first_day, next_day)
        by_day = {row["report_date"]: row for row in rows}
        first = by_day[first_day.isoformat()]
        second = by_day[next_day.isoformat()]
        assert (first["cowork_bill"], first["cafe_bill"], first["discounts_total"],
                first["total_revenue"], first["total_refunds"]) == (100, 415, 28, 487, 0)
        assert (first["cowork_discounts"], first["cafe_discounts"]) == (10, 18)
        assert first["credit_received_methods"]["gcash"] == 60
        assert first["total_orders"] == 1  # The unpaid order did not add to sales.
        assert (second["cowork_bill"], second["cafe_bill"], second["discounts_total"],
                second["total_revenue"], second["total_refunds"], second["total_collections"]) == (0, 60, 0, 60, 100, 150)
        assert second["credit_applied_methods"]["cash"] == 60
        for row in rows:
            assert row["cowork_bill"] + row["cafe_bill"] - row["discounts_total"] == row["total_revenue"]
            assert row["sales_reconciliation_difference"] == 0
            for method, amount in row["checkout_methods"].items():
                assert row["cowork_methods"][method] + row["cafe_methods"][method] - row["discount_methods"][method] == amount
        assert first["checkout_methods"] == {
            "cash": 289, "gcash": 48, "bdo": 100, "bpi": 20, "queenbank": 30, "unclassified": 0,
        }

        generated = SalesService(SalesRepository()).generate_report(first_day, 1, "Audit")
        assert (generated["data"]["cowork_bill"], generated["data"]["cafe_bill"],
                generated["data"]["discounts_total"], generated["data"]["total_refunds"]) == (100, 415, 28, 0)
        assert SalesService(SalesRepository()).list_reports(first_day, first_day)[0]["cowork_bill"] == 100
        context = DailyBalanceExportService(db).build_pdf_context(rows, [])
        assert (context["cowork_bill"], context["cafe_bill"], context["discounts_total"],
                context["total_revenue"], context["total_refunds"]) == (100, 475, 28, 547, 100)
        assert "Checkout Sales Breakdown" in render_template(DailyBalanceExportService.PDF_TEMPLATE, **context)

    client = app.test_client()
    with client.session_transaction() as auth:
        auth.update(user_id=1, username="test_user", role="admin",
                    last_activity=datetime.now(timezone.utc).timestamp())
    query = "start_date=2026-10-05&end_date=2026-10-06"
    api = client.get(f"/admin/daily-balance/api/reports?{query}")
    assert api.status_code == 200
    assert api.get_json()["data"][0]["cafe_bill"] == 60
    csv_response = client.get(f"/admin/daily-balance/api/reports/export-csv?{query}")
    assert csv_response.status_code == 200
    assert "Cowork Bill,Cafe Bill,Discounts,Checkout Sales,Refunds Paid" in csv_response.data.decode("utf-8")
    assert "₱100.00,₱415.00,₱28.00,₱487.00,₱0.00" in csv_response.data.decode("utf-8")
    excel_response = client.get(f"/admin/daily-balance/api/reports/export-excel?{query}")
    assert excel_response.status_code == 200
    workbook = load_workbook(BytesIO(excel_response.data), read_only=True, data_only=True)
    assert workbook["Summary"]["B24"].value == 100
    assert workbook["Summary"]["B25"].value == 475
    assert [workbook["Daily Reports"].cell(row=1, column=column).value for column in range(24, 29)] == [
        "Cowork Bill", "Cafe Bill", "Discounts", "Refunds Paid", "Sales Reconciliation Difference",
    ]
    assert [workbook["Daily Reports"].cell(row=3, column=column).value for column in range(24, 28)] == [
        100, 415, 28, 0,
    ]
    pdf_response = client.get(f"/admin/daily-balance/api/reports/export-pdf?{query}")
    assert pdf_response.status_code == 200
    assert pdf_response.data.startswith(b"%PDF")


def test_legacy_checkout_component_mismatch_is_reported_without_changing_sales(app):
    day = date(2026, 10, 5)
    with app.app_context():
        space = SpaceType(name="Regular Lounge", rate_per_minute=Decimal("0.1667"))
        db.session.add(space)
        db.session.flush()
        session = CustomerSession(customer_name="Legacy", space_type_id=space.id, status="completed")
        db.session.add(session)
        db.session.flush()
        db.session.add(Transaction(session_id=session.id, time_bill=Decimal("20"), food_bill=Decimal("80"),
                                   discount_amount=Decimal("0"), total_bill=Decimal("95"),
                                   payment_method="cash", created_at=datetime(2026, 10, 5, 8)))
        db.session.commit()
        row = SalesRepository().daily_ledger(day, day)[day]
        assert row["total_revenue"] == row["checkout_methods"]["cash"] == 95
        assert (row["cowork_bill"], row["cafe_bill"], row["sales_reconciliation_difference"]) == (20, 80, 5)


def test_additive_upgrade_accepts_legacy_sqlite_schema():
    engine = create_engine("sqlite://")
    with engine.begin() as connection:
        for table in ("users", "payables", "receivable_payments", "expenses",
                      "finance_transactions", "daily_sales_reports", "soft_balance_entries"):
            connection.execute(text(f"CREATE TABLE {table} (id INTEGER PRIMARY KEY)"))
    upgrade(engine)
    upgrade(engine)
    inspector = inspect(engine)
    assert inspector.has_table("payable_payments")
    assert "request_key" in {column["name"] for column in inspector.get_columns("expenses")}
    assert "funding_source" in {column["name"] for column in inspector.get_columns("expenses")}
    assert "funding_source" in {column["name"] for column in inspector.get_columns("payable_payments")}
    assert "total_other_income" in {column["name"] for column in inspector.get_columns("daily_sales_reports")}
    assert "uq_expenses_request_key" in {index["name"] for index in inspector.get_indexes("expenses")}
